#!/usr/bin/env python3
"""Build the editor-facing book proposal (docx zip + html preview)."""
from __future__ import annotations

import zipfile
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "writing" / "drafts" / "book-proposal"
STEM = "出版企画書"
DOCX_NAME = f"{STEM}.docx"
ZIP_NAME = f"{STEM}.zip"

NAVY = RGBColor(0x1A, 0x36, 0x5D)
INK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
RULE = RGBColor(0xC4, 0x5C, 0x26)

# kind: kicker | title | catch | meta | h1 | h2 | p | li | note
BLOCKS: list[tuple[str, str]] = [
    ("kicker", "出版企画書"),
    ("title", "『昭和というロストテクノロジー』"),
    ("catch", "仕事のコミュニケーションコストは下げろ。人間のコミュニケーションコストは上げろ。"),
    ("meta", "野田浩平（主著）／飯田史也／猪原健弘／松岡良彦"),
    ("meta", "2026年9月"),
    ("h1", "1. タイトル"),
    (
        "p",
        "いま表紙に置くなら、これです。棚で手に取ってもらうための入口であって、中身の要約ではありません。昭和の全部を褒める本でもありません。捨ててしまった会うこと、雑談、世話、失敗の共有を、もう一度設計し直す話です。題は、棚に合わせて変えても同じ本です。",
    ),
    ("h2", "いまの案"),
    ("li", "『昭和というロストテクノロジー』"),
    ("li", "副題：捨てたのはパワハラではない。会う、雑談、世話、失敗の共有である"),
    ("li", "帯：仕事のコミュニケーションコストは下げろ。人間のコミュニケーションコストは上げろ。"),
    (
        "note",
        "戻さないのは、パワハラ、飲酒の強要、理不尽な上下です。残すのは、身体を同じ場所に置くことです。",
    ),
    ("h2", "ほかの候補"),
    ("li", "仕事のコミュニケーションコストは下げろ。人間のコミュニケーションコストは上げろ。（題にもできる）"),
    ("li", "『AIに代替されない「最後の知性」』"),
    ("li", "『「正解のない世界」で、自分を失わない技術』"),
    ("li", "『「違和感」を無視して生きてきた大人たちへ』"),
    ("h1", "2. 端的にいうと"),
    (
        "p",
        "仕事は速くなった。人が動かない。頭の作業はAIに渡せる。残るのは、会うことと、身体と、人と場をつくる力だ。",
    ),
    ("h2", "書店で手に取る理由"),
    ("li", "仕事のコストは下げろ。人間のコストは上げろ。"),
    ("li", "ChatGPTに聞くな。会いに行け。"),
    ("li", "昭和の全部を戻せ、ではない。会う、雑談、世話、失敗の共有を、設計し直せ。"),
    ("h1", "3. なぜ今か"),
    (
        "p",
        "資料も議事録もコードも、速く安くなった。それでも人は動かない。学校では特別活動が、人と決める訓練場になっている。家では、STEMは足りている、次は何か、という親がいる。会社では、仕事は回るが、現場の空気が読めない、という声がある。",
    ),
    (
        "p",
        "身体、人と場、内面。こういう力は、経営教育でも、教育学でも、エリート教育の見直しでも、以前から言われてきた。特定の大学や国際機関の名前を出せば新しく見えるが、新しさはそこではない。この本が書くのは、AIのあとに、日本の家庭、学校、職場で何を育てるかだ。",
    ),
    (
        "p",
        "診断の本ではありません。学びと現場の本です。",
    ),
    ("h1", "4. 読者"),
    (
        "p",
        "中心は、親でもある大人です。",
    ),
    ("li", "STEMに関心のある親。知識は与えた。次に何を育てるか。"),
    ("li", "教師。特別活動、学級、探究。人と合意の作り方。"),
    ("li", "経営者、マネジャー。人が離れない現場。"),
    (
        "p",
        "若手には「昭和」へのアレルギーがあるので、昭和は入口に使い、本文は現場の話で読ませます。題は変えても構いません。",
    ),
    ("h1", "5. 中心概念"),
    (
        "p",
        "頭のよさだけでは足りない。身体、人と場、言葉、内面。そのあたりを、この顔ぶれで書く。",
    ),
    (
        "p",
        "飯田は身体とものづくり（身体性ロボット、人工知能）。猪原は人と集団で決めること（合意形成、特別活動）。野田は認知科学と学びの設計。松岡は言語と現場の学び（RLE、子どもの英語教育）。",
    ),
    ("h2", "現場で使う地図"),
    ("li", "身体を読む——空腹、疲労、呼吸。"),
    ("li", "場を読む——表情、空気、期待。"),
    ("li", "言葉にする——攻撃でも沈黙でもなく。"),
    ("li", "余白をつくる——衝動と行動のあいだ。"),
    ("li", "共に決める——違うまま、乗せる。"),
    ("h2", "学術的な背景（五つの層）"),
    (
        "p",
        "本文の題や帯には出しません。裏にある整理です。身体、予測、意味、関係、集団・社会。教育や認知の議論では、多重知能のように、頭以外の力を分けて見る見方もあります。看板にはしません。",
    ),
    ("li", "身体——心拍、呼吸、空腹、疲労。"),
    ("li", "予測——予想と現実のずれ。驚き、学び。"),
    ("li", "意味——何を大切にするか。"),
    ("li", "関係——表情、態度、期待。"),
    ("li", "集団・社会——同調、制度、合意形成。特別活動。"),
    ("h1", "6. 章構成"),
    (
        "p",
        "序章、五章、終章。題は現場の言葉です。誰がどの章を書くかは、これから組みます。第3章は、いま決めません。",
    ),
    ("h2", "序章　捨ててしまったのは、会うことである"),
    ("li", "なぜ「プログラミングができる子ども」ほど、先が不安なのか"),
    ("li", "仕事は速くなった。人が動かない。"),
    ("li", "会うことを、設計し直す"),
    ("h2", "第1章　頭がいい人ほど、先に折れる"),
    ("li", "ChatGPTに聞くな。会いに行け"),
    ("li", "データは手段を選ぶ。何を大切にするかは、人が決める"),
    ("li", "正しいことが、人を動かさない理由"),
    ("h2", "第2章　AIは、お腹が空かない"),
    ("li", "ロボット研究者が、胃腸を見る理由"),
    ("li", "画面の前の論理が、こぼしているもの"),
    ("li", "予測が外れたとき、人は学び始める"),
    ("h2", "第3章　0.5秒の余白をつくる"),
    ("li", "空腹を、性格だと思わない"),
    ("li", "衝動と行動のあいだに、選択肢を置く"),
    ("li", "沈黙と爆発は、同じ罠の両側である"),
    ("h2", "第4章　リーダーの一声で、場の頭が悪くなる"),
    ("li", "「何でも言っていい」は、安全ではない"),
    ("li", "会う、雑談、世話、失敗の共有"),
    ("li", "場は、設計できる"),
    ("h2", "第5章　正しいことを言っても、人は動かない"),
    ("li", "特別活動は、日本がすでに持っている訓練場である"),
    ("li", "SNSは、理性を手放させる"),
    ("li", "対立の奥にある恐怖・屈辱・意地"),
    ("h2", "終章　仕事のコストは下げろ。人間のコストは上げろ。"),
    ("li", "親が先に、自分の身体を読む"),
    ("li", "学校は、共に決める場をつくる"),
    ("li", "経営は、会いに行く"),
    (
        "note",
        "各章の終わりに、短い実践を一つ置きます。身体の天気予報、言葉の翻訳、驚きノート、場の地図、対話の問い、など。",
    ),
    ("h1", "7. 著者"),
    (
        "p",
        "主著は野田。飯田、猪原、松岡が加わる。松岡は野田のビジネスと研究のパートナーで、著者の末尾に置きます。",
    ),
    ("h2", "飯田史也"),
    (
        "p",
        "ケンブリッジ大学工学部ロボティクス教授。東京大学大学院工学系研究科教授（精密工学）。身体性ロボット、身体性AI、バイオインスパイアード・ロボティクス。BIRLをケンブリッジと東京で主宰。関連書に『世界最高峰の学び』。",
    ),
    ("h2", "猪原健弘"),
    (
        "p",
        "東京科学大学リベラルアーツ研究教育院教授。中央教育審議会 初等中等教育分科会 教育課程部会 特別活動ワーキンググループ専門委員（2025年9月〜2027年3月）。社会工学、意思決定、ゲーム理論、合意形成。",
    ),
    ("h2", "野田浩平（主著）"),
    (
        "p",
        "認知科学者（博士）。株式会社ココロラボ代表取締役。グロービス経営大学院専任教員。MIT IDEAS Asia Pacific Regional Faculty。認知科学とリーダーシップ教育。",
    ),
    ("h2", "松岡良彦（ビジネス・研究パートナー）"),
    (
        "p",
        "野田のビジネスと研究のパートナー。株式会社Ducks and Drakes。フィリピン・ドゥマゲテの Starting Point English Academy（SPEA）で、子どもの英語教育を現場でともに見ていた。2016年、日本認知科学会第33回大会で野田と共著発表（第二言語としての英語学習における Real Life Experience 法の提案）。",
    ),
    ("h1", "8. 仕様"),
    (
        "p",
        "判型、ページ数、価格は、いま決めません。厚い本に限りません。書店で手に取ってもらえる短さも含めて、編集者と相談したいです。公開講座や鼎談は、本のあとの話です。",
    ),
]


def set_run(run, *, size=11, bold=False, color=INK, name="Yu Mincho"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def add_p(doc, text, *, size=11, bold=False, color=INK, space_after=8, space_before=0, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.35
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color)
    return p


def build_docx(path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    for kind, text in BLOCKS:
        if kind == "kicker":
            add_p(doc, text, size=12, bold=True, color=MUTED, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif kind == "title":
            add_p(doc, text, size=18, bold=True, color=NAVY, space_after=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif kind == "catch":
            add_p(doc, text, size=12, bold=True, color=RULE, space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif kind == "meta":
            add_p(doc, text, size=10, color=MUTED, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif kind == "h1":
            add_p(doc, text, size=15, bold=True, color=NAVY, space_before=16, space_after=8)
        elif kind == "h2":
            add_p(doc, text, size=12, bold=True, color=NAVY, space_before=10, space_after=6)
        elif kind == "p":
            p = add_p(doc, text, size=11, space_after=8)
            p.paragraph_format.first_line_indent = Cm(0.7)
        elif kind == "li":
            p = add_p(doc, "・" + text, size=11, space_after=3)
            p.paragraph_format.left_indent = Cm(0.7)
        elif kind == "note":
            add_p(doc, text, size=10, color=MUTED, space_before=4, space_after=8)
        else:
            raise ValueError(kind)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def build_md(path: Path) -> None:
    lines: list[str] = []
    prev = ""
    for kind, text in BLOCKS:
        if prev == "li" and kind != "li":
            lines.append("")
        if kind == "kicker":
            lines.append(f"*{text}*")
            lines.append("")
        elif kind == "title":
            lines.append(f"# {text}")
            lines.append("")
        elif kind == "catch":
            lines.append(f"**{text}**")
            lines.append("")
        elif kind == "meta":
            lines.append(text)
            lines.append("")
        elif kind == "h1":
            lines.append(f"## {text}")
            lines.append("")
        elif kind == "h2":
            lines.append(f"### {text}")
            lines.append("")
        elif kind == "p":
            lines.append(text)
            lines.append("")
        elif kind == "li":
            lines.append(f"- {text}")
        elif kind == "note":
            lines.append(f"> {text}")
            lines.append("")
        else:
            raise ValueError(kind)
        prev = kind
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


RAW_ZIP = (
    "https://github.com/kohnoda-glitch/kocorolab/raw/cursor/note-medium-writing-4caf"
    "/writing/drafts/book-proposal/出版企画書.zip"
)


def build_html(path: Path) -> None:
    parts = [
        "<!DOCTYPE html>",
        '<html lang="ja">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        f"<title>{escape(STEM)}</title>",
        "<style>",
        "body { font-family: 'Yu Mincho', 'Hiragino Mincho ProN', serif; max-width: 720px; margin: 24px auto; padding: 0 16px 48px; line-height: 1.75; color: #222; }",
        ".note { font-family: sans-serif; background: #fff6d8; border: 1px solid #e6d48a; padding: 12px 16px; margin-bottom: 24px; font-size: 0.95rem; }",
        "h1 { font-size: 1.55rem; color: #1a365d; text-align: center; line-height: 1.4; }",
        "h2 { font-size: 1.2rem; color: #1a365d; margin-top: 2em; }",
        "h3 { font-size: 1.05rem; color: #1a365d; margin-top: 1.4em; }",
        ".kicker, .meta { text-align: center; color: #555; font-size: 0.95rem; }",
        ".catch { text-align: center; color: #c45c26; font-weight: bold; }",
        "ul { padding-left: 1.2em; }",
        "aside { color: #555; font-size: 0.95rem; border-left: 3px solid #c45c26; padding-left: 12px; margin: 1em 0; }",
        "@media print { .note { display: none; } body { max-width: none; } }",
        "</style>",
        "</head>",
        "<body>",
        '<div class="note">',
        "編集者に渡すのは Word です。Chrome では下のリンクを右クリック → 名前を付けてリンク先を保存。",
        f'<br><a href="{escape(RAW_ZIP)}">{escape(ZIP_NAME)}</a>',
        "</div>",
        "<article>",
    ]
    open_ul = False

    def close_ul():
        nonlocal open_ul
        if open_ul:
            parts.append("</ul>")
            open_ul = False

    for kind, text in BLOCKS:
        t = escape(text)
        if kind != "li":
            close_ul()
        if kind == "kicker":
            parts.append(f'<p class="kicker">{t}</p>')
        elif kind == "title":
            parts.append(f"<h1>{t}</h1>")
        elif kind == "catch":
            parts.append(f'<p class="catch">{t}</p>')
        elif kind == "meta":
            parts.append(f'<p class="meta">{t}</p>')
        elif kind == "h1":
            parts.append(f"<h2>{t}</h2>")
        elif kind == "h2":
            parts.append(f"<h3>{t}</h3>")
        elif kind == "p":
            parts.append(f"<p>{t}</p>")
        elif kind == "li":
            if not open_ul:
                parts.append("<ul>")
                open_ul = True
            parts.append(f"<li>{t}</li>")
        elif kind == "note":
            parts.append(f"<aside>{t}</aside>")
        else:
            raise ValueError(kind)
    close_ul()
    parts.extend(["</article>", "</body>", "</html>", ""])
    path.write_text("\n".join(parts), encoding="utf-8")


def build_zip(docx: Path, zpath: Path) -> None:
    with zipfile.ZipFile(zpath, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.write(docx, DOCX_NAME)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    md = OUT / f"{STEM}.md"
    docx = OUT / DOCX_NAME
    zpath = OUT / ZIP_NAME
    html = OUT / "OPEN-IN-BROWSER.html"
    build_md(md)
    build_docx(docx)
    build_zip(docx, zpath)
    build_html(html)
    print(docx)
    print(zpath)
    print(html)


if __name__ == "__main__":
    main()
